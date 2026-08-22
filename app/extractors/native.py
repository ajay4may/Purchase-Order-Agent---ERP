"""Native Office, spreadsheet, CSV, and Outlook message extraction."""

from __future__ import annotations

import csv
import io
import tempfile
import zipfile
from email import policy
from email.parser import BytesParser
from pathlib import Path

import extract_msg
import openpyxl
import xlrd
from docx import Document
from extract_msg.exceptions import InvalidFileFormatError
from openpyxl.utils import get_column_letter

from app.errors import CorruptDocumentError, EncryptedDocumentError, UnsupportedFormatError
from app.models import Evidence, RawDocument


class NativeDocumentExtractor:
    def __init__(
        self,
        *,
        max_uncompressed_bytes: int = 200 * 1024 * 1024,
        max_archive_entries: int = 5000,
    ) -> None:
        self._max_uncompressed_bytes = max_uncompressed_bytes
        self._max_archive_entries = max_archive_entries

    async def extract(self, file_name: str, mime_type: str, content: bytes) -> RawDocument:
        extension = Path(file_name).suffix.lower()
        try:
            if extension == ".docx" or "wordprocessingml" in mime_type:
                return self._docx(content)
            if extension == ".xlsx" or "spreadsheetml" in mime_type:
                return self._xlsx(content)
            if extension == ".xls" or mime_type == "application/vnd.ms-excel":
                return self._xls(content)
            if extension == ".csv" or mime_type in {"text/csv", "application/csv"}:
                return self._csv(content)
            if extension == ".msg" or mime_type == "application/vnd.ms-outlook":
                return self._msg(content)
            if extension == ".eml" or mime_type == "message/rfc822":
                return self._eml(content)
            if mime_type.startswith("text/"):
                return self._text(content)
        except EncryptedDocumentError:
            raise
        except (ValueError, KeyError, OSError, UnicodeError, zipfile.BadZipFile) as exc:
            raise CorruptDocumentError(f"Could not parse {file_name}.") from exc
        raise UnsupportedFormatError(f"Unsupported native document format: {mime_type}")

    def _docx(self, content: bytes) -> RawDocument:
        self._validate_zip_container(content)
        try:
            document = Document(io.BytesIO(content))
        except (ValueError, OSError, zipfile.BadZipFile) as exc:
            if "password" in str(exc).lower() or "encrypted" in str(exc).lower():
                raise EncryptedDocumentError("Encrypted DOCX files are not supported.") from exc
            raise
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        tables: list[list[list[str | None]]] = []
        for table in document.tables:
            tables.append([[cell.text for cell in row.cells] for row in table.rows])
        table_text = NativeDocumentExtractor._tables_to_text(tables)
        return RawDocument(
            text="\n".join(paragraphs + table_text),
            tables=tables,
            extracted_with="python-docx",
            evidence=[Evidence(source_type="document", text=text) for text in paragraphs],
        )

    def _xlsx(self, content: bytes) -> RawDocument:
        self._validate_zip_container(content)
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (ValueError, OSError) as exc:
            if "encrypted" in str(exc).lower():
                raise EncryptedDocumentError("Encrypted XLSX files are not supported.") from exc
            raise
        tables: list[list[list[str | None]]] = []
        evidence: list[Evidence] = []
        lines: list[str] = []
        for sheet in workbook.worksheets:
            matrix: list[list[str | None]] = []
            lines.append(f"[sheet {sheet.title}]")
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [None if value is None else str(value) for value in row]
                if not any(value not in {None, ""} for value in values):
                    continue
                matrix.append(values)
                lines.append("\t".join(value or "" for value in values))
                for column_index, value in enumerate(values, start=1):
                    if value:
                        evidence.append(
                            Evidence(
                                source_type="spreadsheet",
                                text=value,
                                sheet=sheet.title,
                                cell=f"{get_column_letter(column_index)}{row_index}",
                            )
                        )
            if matrix:
                tables.append(matrix)
        return RawDocument(
            text="\n".join(lines),
            tables=tables,
            sheet_count=len(workbook.worksheets),
            extracted_with="openpyxl",
            evidence=evidence,
        )

    @staticmethod
    def _xls(content: bytes) -> RawDocument:
        try:
            workbook = xlrd.open_workbook(file_contents=content)
        except xlrd.biffh.XLRDError as exc:
            if "encrypted" in str(exc).lower() or "password" in str(exc).lower():
                raise EncryptedDocumentError("Encrypted XLS files are not supported.") from exc
            raise
        tables: list[list[list[str | None]]] = []
        evidence: list[Evidence] = []
        lines: list[str] = []
        for sheet in workbook.sheets():
            matrix: list[list[str | None]] = []
            lines.append(f"[sheet {sheet.name}]")
            for row_index in range(sheet.nrows):
                values = [
                    str(value) if value != "" else None for value in sheet.row_values(row_index)
                ]
                if not any(values):
                    continue
                matrix.append(values)
                lines.append("\t".join(value or "" for value in values))
                for column_index, value in enumerate(values):
                    if value:
                        evidence.append(
                            Evidence(
                                source_type="spreadsheet",
                                text=value,
                                sheet=sheet.name,
                                cell=f"{get_column_letter(column_index + 1)}{row_index + 1}",
                            )
                        )
            if matrix:
                tables.append(matrix)
        return RawDocument(
            text="\n".join(lines),
            tables=tables,
            sheet_count=workbook.nsheets,
            extracted_with="xlrd",
            evidence=evidence,
        )

    @staticmethod
    def _csv(content: bytes) -> RawDocument:
        text = content.decode("utf-8-sig")
        rows: list[list[str | None]] = [
            [value for value in row] for row in csv.reader(io.StringIO(text))
        ]
        evidence: list[Evidence] = []
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                if value:
                    evidence.append(
                        Evidence(
                            source_type="spreadsheet",
                            text=value,
                            sheet="CSV",
                            cell=f"{get_column_letter(column_index)}{row_index}",
                        )
                    )
        return RawDocument(
            text="\n".join("\t".join(value or "" for value in row) for row in rows),
            tables=[rows],
            sheet_count=1,
            extracted_with="python-csv",
            evidence=evidence,
        )

    @staticmethod
    def _msg(content: bytes) -> RawDocument:
        path: str | None = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".msg", delete=False) as handle:
                handle.write(content)
                path = handle.name
            message = extract_msg.Message(path)  # type: ignore[no-untyped-call]
            try:
                body = message.body or ""
                subject = message.subject or ""
                lines = [f"Subject: {subject}", str(body)]
                attachments: list[str] = []
                for attachment in message.attachments:
                    name = getattr(attachment, "longFilename", None) or getattr(
                        attachment, "shortFilename", None
                    )
                    if name:
                        attachments.append(str(name))
                    data = getattr(attachment, "data", None)
                    if (
                        isinstance(data, bytes)
                        and name
                        and Path(name).suffix.lower()
                        in {
                            ".txt",
                            ".csv",
                        }
                    ):
                        attachments.append(data.decode("utf-8", errors="replace"))
                if attachments:
                    lines.append("[attachments]\n" + "\n".join(attachments))
                text = "\n".join(lines)
                return RawDocument(
                    text=text,
                    extracted_with="extract-msg",
                    evidence=[Evidence(source_type="email", text=text)],
                )
            finally:
                message.close()
        except InvalidFileFormatError as exc:
            raise CorruptDocumentError("Could not parse the Outlook MSG file.") from exc
        except (AttributeError, OSError, ValueError) as exc:
            if "encrypted" in str(exc).lower():
                raise EncryptedDocumentError(
                    "Encrypted Outlook MSG files are not supported."
                ) from exc
            raise
        finally:
            if path:
                Path(path).unlink(missing_ok=True)

    @staticmethod
    def _eml(content: bytes) -> RawDocument:
        message = BytesParser(policy=policy.default).parsebytes(content)
        parts: list[str] = [f"Subject: {message.get('subject', '')}"]
        if message.is_multipart():
            for part in message.walk():
                if part.get_content_type() == "text/plain":
                    parts.append(part.get_content())
        else:
            parts.append(message.get_content())
        text = "\n".join(parts)
        return RawDocument(
            text=text,
            extracted_with="python-email",
            evidence=[Evidence(source_type="email", text=text)],
        )

    @staticmethod
    def _text(content: bytes) -> RawDocument:
        text = content.decode("utf-8-sig")
        return RawDocument(
            text=text,
            extracted_with="text-decoder",
            evidence=[Evidence(source_type="email", text=text)],
        )

    @staticmethod
    def _tables_to_text(tables: list[list[list[str | None]]]) -> list[str]:
        return [
            "\n".join("\t".join(cell or "" for cell in row) for row in table) for table in tables
        ]

    def _validate_zip_container(self, content: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > self._max_archive_entries:
                    raise CorruptDocumentError("Office document contains too many archive entries.")
                total_uncompressed = sum(entry.file_size for entry in entries)
                if total_uncompressed > self._max_uncompressed_bytes:
                    raise CorruptDocumentError(
                        "Office document expands beyond the safe processing limit."
                    )
                total_compressed = sum(max(entry.compress_size, 1) for entry in entries)
                if total_uncompressed > 10 * 1024 * 1024 and (
                    total_uncompressed / total_compressed > 1000
                ):
                    raise CorruptDocumentError("Office document has an unsafe compression ratio.")
        except zipfile.BadZipFile as exc:
            raise CorruptDocumentError("Office document is not a valid ZIP container.") from exc
